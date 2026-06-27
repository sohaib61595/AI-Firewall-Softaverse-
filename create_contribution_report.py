import os
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

title = doc.add_heading('AI Firewall: Contribution Report', level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph('This document outlines the specific roles and contributions of each team member towards the successful development of the AI Firewall: LLM Prompt Injection Scanner project.')

doc.add_heading('Team Contributions', level=1)

contributions = [
    ("Shayan Haider", "Frontend UI/UX Developer", "Developed the interactive web interface where users test prompts. Ensured clear visual feedback and seamless dashboard interactions."),
    ("Muhammad Sohaib Nisar", "NLP & Machine Learning Developer", "Handled the core text classification logic. Cleaned raw text data, implemented tokenization/vectorization, and trained a high-accuracy Scikit-learn classifier."),
    ("Chaudhary Abdul Rafay", "Backend API Developer", "Built the FastAPI routing connecting the frontend to the NLP model, ensuring low-latency data transmission and robust API endpoints."),
    ("Syed Zohaib Abbas", "Database & Logging Architect", "Set up a secure SQLite database to capture attack data. Built the backend logic to reliably log every blocked prompt for future auditing."),
    ("Shahryar Ahmed", "QA, Red Teaming & Documentation", "Managed all SQA testing logs, rigorously tested the system against bypasses, maintained project documentation, and handled final presentation formatting.")
]

table = doc.add_table(rows=1, cols=3)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = 'Team Member'
hdr_cells[1].text = 'Role'
hdr_cells[2].text = 'Specific Contributions'

for member, role, desc in contributions:
    row_cells = table.add_row().cells
    row_cells[0].text = member
    row_cells[1].text = role
    row_cells[2].text = desc

doc.add_paragraph('\n')
doc.add_paragraph('All members successfully collaborated to ensure the system architecture, documentation, and source code met the required professional standards.')

doc.save('Contribution_Report.docx')
print("Successfully created Contribution_Report.docx")
